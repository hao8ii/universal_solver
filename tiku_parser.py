# -*- coding: utf-8 -*-
"""题库格式转换器

把 Word（.docx/.doc/.wps）和 PDF 题库转换成“云端学习答题工具”可直接读取的
Excel 文件。输出 Excel 的前两列固定为“题目”“答案”，其余列仅供查看。

本文件同时作为“一体化版”的内部解析模块（import_tiku_file）。
也可以直接运行，作为独立的 Word/PDF → Excel 转换器：

    python tiku_importer.py                     # 打开图形界面
    python tiku_importer.py 题库.docx           # 生成 题库_题库.xlsx
    python tiku_importer.py a.docx b.pdf --merge -o 合集.xlsx
    python tiku_importer.py 题库.pdf --preview 20

依赖：python-docx、pypdf、openpyxl；.doc/.wps 另需 Windows 上的
Microsoft Word 或 WPS（通过 COM 自动转换）。
"""

import argparse
import os
import re
import shutil
import sys
import tempfile
import threading
import traceback
import zipfile
from collections import Counter

APP_NAME = "题库格式转换器"
VERSION = "1.0"

SUPPORTED_EXTS = (".docx", ".doc", ".wps", ".pdf")

# 题号：1、 / 1. / 1．
QSTART = re.compile(r"^\s*(\d{1,5})\s*[、.．]\s*(.*)$")
# 选项：A. / A、 / Ａ．
OPTSTART = re.compile(r"^\s*([A-GＡ-Ｇ])\s*[、.．]\s*(.*)$")
# “正确答案：B”“参考答案：A、B”
ANS_LINE = re.compile(
    r"(?:正确|参考)?答案\s*[：:]\s*"
    r"([A-Ga-g](?:\s*[,，、;；]?\s*[A-Ga-g]){0,7}|正确|错误)"
)
# 判断题答案：(√) (×)
JUDGE_MARK = re.compile(r"[（(]\s*([√×✓✗])\s*[）)]")
# 括号中的选项字母：(A) (ABCD)
PAREN_LETTERS = re.compile(r"[（(]\s*([A-Ga-g]{1,8})\s*[）)]")
# 一、判断题 / 二、单选题
SECTION_RE = re.compile(r"^\s*[一二三四五六七八九十]+\s*[、.．]\s*(.{1,24})$")
# 题干末尾只有选项标号、没有内容的占位符：A.  B.  C.  D.
TRAILING_LABELS = re.compile(r"(?:\s*[A-GＡ-Ｇ]\s*[.、．]\s*){2,}$")
LETTERS_RE = re.compile(r"^[A-G]{1,8}$")
# 行尾单独出现的答案字母（前面必须是空白或中文标点）
TRAILING_ANSWER = re.compile(
    r"(?:[。？?！!；;：:）)]|\s)([A-Ga-g]{1,8})\s*[。.]?\s*$"
)
# 中文之间的多余空格
CJK_SPACE = re.compile(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])")
# Excel 不允许的控制字符
EXCEL_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class ConvertError(Exception):
    pass


# ---------------------------------------------------------------------------
# 文本工具
# ---------------------------------------------------------------------------

def _append_note(note, text):
    if not text:
        return note or ""
    if not note:
        return text
    if text in note:
        return note
    return note + "；" + text


def join_wrapped(lines):
    """把 Word/PDF 自动换行产生的多行合并成一行。"""
    out = ""
    for raw in lines:
        s = (raw or "").strip()
        if not s:
            continue
        if out and re.search(r"[A-Za-z0-9]$", out) and re.match(r"[A-Za-z0-9]", s):
            out += " "
        out += s
    return out.strip()


def strip_number_prefix(text):
    # 有些文档会出现“218.10.根据……”这种重复编号
    return re.sub(r"^\s*(?:\d{1,5}\s*[、.．]\s*)+", "", text or "").strip()


def strip_trailing_labels(text):
    return TRAILING_LABELS.sub("", text or "").strip()


def normalize_spaces(text):
    text = (text or "").replace("\u3000", " ").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = CJK_SPACE.sub("", text)
    return text.strip()


def clean_stem(text):
    text = strip_number_prefix(text)
    text = strip_trailing_labels(text)
    text = normalize_spaces(text)
    text = re.sub(r"^[\s、，,。.；;：:]+", "", text)
    return text.strip()


def normalize_letters(text):
    letters = re.findall(r"[A-G]", (text or "").upper())
    return "".join(dict.fromkeys(letters))


def normalize_section(text):
    if not text:
        return None
    if "判断" in text:
        return "判断题"
    if "多" in text and "选" in text:
        return "多选题"
    if "单" in text and "选" in text:
        return "单选题"
    if "选择" in text:
        return "单选题"
    return None


def is_section_header(text):
    return any(k in text for k in
               ("判断题", "单选题", "多选题", "单项选择题", "多项选择题"))


def _is_judgment_options(opt_map):
    """A 选项是“正确”、B 选项是“错误”时，判定为判断题。"""
    if not opt_map or set(opt_map) != set("AB"):
        return False
    a = re.sub(r"[\s。.、，,；;：:（）()]", "", opt_map.get("A", ""))
    b = re.sub(r"[\s。.、，,；;：:（）()]", "", opt_map.get("B", ""))
    a_true = "正确" in a or a in ("对", "√", "是")
    b_false = "错误" in b or b in ("错", "×", "否")
    return a_true and b_false


def extract_options(lines):
    """从选项行中提取 {字母: 选项内容}，兼容一行多个选项。"""
    text = " ".join((x or "").strip() for x in lines if (x or "").strip())
    if not text:
        return {}
    matches = list(re.finditer(r"(?:^|\s)([A-GＡ-Ｇ])\s*[、.．]\s*", text))
    result = {}
    for i, m in enumerate(matches):
        letter = m.group(1)
        letter = letter.translate(str.maketrans("ＡＢＣＤＥＦＧ", "ABCDEFG"))
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        value = text[m.end():end].strip()
        if letter not in result:
            result[letter] = value
    return result


def classify_answer(answer, section=None, judgment=False):
    if judgment or answer in ("正确", "错误"):
        return "判断题"
    if answer and LETTERS_RE.match(answer):
        return "多选题" if len(answer) > 1 else "单选题"
    if section:
        return section
    return "其他"


def _marker_replacement(text, match):
    """句末的答案标记直接删掉，句中的答案标记替换成空括号。"""
    tail = text[match.end():]
    if re.fullmatch(r"[\s。.、，,；;：:！!？?]*", tail):
        return text[:match.start()] + text[match.end():]
    return text[:match.start()] + "（ ）" + text[match.end():]


def extract_inline_answer(stem):
    """从题干内提取答案，返回 (答案, 清洗后的题干, 是否判断题)。"""
    m = JUDGE_MARK.search(stem)
    if m:
        answer = "正确" if m.group(1) in "√✓" else "错误"
        return answer, _marker_replacement(stem, m), True

    matches = list(PAREN_LETTERS.finditer(stem))
    if matches:
        letters = "".join(m.group(1).upper() for m in matches)
        answer = "".join(dict.fromkeys(letters))
        last = matches[-1]
        tail = stem[last.end():]
        if re.fullmatch(r"[\s。.、，,；;：:！!？?]*", tail):
            cleaned = PAREN_LETTERS.sub("", stem)
        else:
            cleaned = PAREN_LETTERS.sub("（ ）", stem)
        return answer, cleaned, False

    return None, stem, False


def find_inline_options_start(text):
    """找题干同一行里真正的选项起点，例如“问题 A. 甲 B. 乙”。"""
    m = re.search(r"\s+[AＡ]\s*[.、．]\s*\S", text)
    if not m:
        return None
    rest = text[m.end():]
    if re.search(r"\s+[B-Ｇ]\s*[.、．]\s*\S", rest):
        return m.start()
    return None


def clean_answer_markers(stem, answer):
    """题干里若把答案字母写进了括号（如“其（ABCD）等情况”），替换成空括号。"""
    if not answer or not LETTERS_RE.match(answer):
        return stem
    target = set(answer)

    def repl(m):
        letters = set(m.group(1).upper())
        if not letters.issubset(target):
            return m.group(0)
        tail = stem[m.end():]
        if re.fullmatch(r"[\s。.、，,；;：:！!？?]*", tail):
            return ""
        return "（ ）"

    return PAREN_LETTERS.sub(repl, stem)


# ---------------------------------------------------------------------------
# 文件读取
# ---------------------------------------------------------------------------

def _read_docx_lines(path):
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise ConvertError("缺少 python-docx，请先执行：pip install python-docx")

    doc = Document(path)
    lines = []

    def add_table(table):
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                key = cell._tc
                if key in seen:
                    continue
                seen.add(key)
                for p in cell.paragraphs:
                    lines.append(p.text)
                for nested in cell.tables:
                    add_table(nested)

    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            lines.append(Paragraph(child, doc).text)
        elif tag == "tbl":
            add_table(Table(child, doc))
    return lines


def _read_pdf_lines(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ConvertError("缺少 pypdf，请先执行：pip install pypdf")

    try:
        reader = PdfReader(path)
        if getattr(reader, "is_encrypted", False):
            try:
                if not reader.decrypt(""):
                    raise ConvertError("PDF 已加密，无法读取")
            except ConvertError:
                raise
            except Exception:
                raise ConvertError("PDF 已加密，无法读取")
        lines = []
        for page in reader.pages:
            text = page.extract_text() or ""
            for raw in text.replace("\r", "\n").split("\n"):
                s = raw.strip()
                if not s:
                    continue
                # 去掉单独的页码
                if re.fullmatch(r"\d{1,4}", s):
                    continue
                lines.append(s)
    except ConvertError:
        raise
    except Exception as e:
        raise ConvertError("读取 PDF 失败：%s" % e)

    if not lines:
        raise ConvertError("PDF 中没有可提取的文字，可能是扫描件；请先 OCR 或另存为可复制文字的 PDF。")
    return lines


def _office_convert_to_docx(src, dst):
    """用 Word/WPS 的 COM 接口把 .doc/.wps 转成 .docx。"""
    try:
        import win32com.client
    except ImportError:
        raise ConvertError(
            "读取 .doc/.wps 需要 pywin32 和已安装的 Word/WPS。\n"
            "可先手动另存为 .docx，或执行：pip install pywin32"
        )

    last_err = None
    for prog_id in ("Word.Application", "KWps.Application", "WPS.Application"):
        app = None
        try:
            app = win32com.client.DispatchEx(prog_id)
            try:
                app.Visible = False
            except Exception:
                pass
            try:
                app.DisplayAlerts = 0
            except Exception:
                pass
            doc = app.Documents.Open(os.path.abspath(src), ReadOnly=True)
            try:
                doc.SaveAs2(os.path.abspath(dst), FileFormat=16)
            except Exception:
                doc.SaveAs(os.path.abspath(dst), FileFormat=16)
            doc.Close(False)
            app.Quit()
            return
        except Exception as e:
            last_err = e
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
    raise ConvertError(
        "无法用 Word/WPS 打开该文件（可能未安装或文件损坏）。\n"
        "请先手动另存为 .docx 后再转换。\n原始错误：%s" % last_err
    )


def read_source_lines(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return _read_pdf_lines(path)
    if ext == ".docx":
        return _read_docx_lines(path)
    if ext in (".doc", ".wps"):
        # 有些 .wps 实际是 docx（zip）格式
        if zipfile.is_zipfile(path):
            try:
                return _read_docx_lines(path)
            except Exception:
                pass
        tmp_dir = tempfile.mkdtemp(prefix="tiku_conv_")
        try:
            tmp_docx = os.path.join(tmp_dir, "converted.docx")
            _office_convert_to_docx(path, tmp_docx)
            return _read_docx_lines(tmp_docx)
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)
    raise ConvertError("不支持的文件格式：%s（支持 %s）" %
                       (ext or "无扩展名", "、".join(SUPPORTED_EXTS)))


# ---------------------------------------------------------------------------
# 解析器
# ---------------------------------------------------------------------------

def _new_question(no, qtype, stem, answer, note="", options=None):
    return {
        "no": no,
        "qtype": qtype,
        "stem": stem,
        "answer": answer or "",
        "note": note or "",
        # 保留 A/B/C/D 选项原文，通用答题程序需要按选项内容匹配，
        # 不能只依赖题库里的字母顺序。
        "options": dict(options or {}),
    }


def _finish_question(no, section, stem, answer, judgment, options, note=""):
    qtype = classify_answer(answer, section, judgment)
    if section == "判断题" and qtype != "判断题" and answer:
        note = _append_note(note, "章节标注为判断题")
    elif section and qtype not in ("判断题", section) and answer:
        note = _append_note(note, "章节为%s但答案为%s" % (section, answer))

    if answer and LETTERS_RE.match(answer) and options:
        missing = [ch for ch in answer if ch not in options]
        if missing:
            note = _append_note(note, "答案字母%s不在已识别选项中" % "".join(missing))

    if not stem:
        note = _append_note(note, "题干为空")
    return _new_question(no, qtype, stem, answer, note, options)


def parse_answerline(lines, source=""):
    """解析“题干 + 选项 + 正确答案：X”格式（2023 Word、2025 PDF）。"""
    blocks = []
    cur = None
    expected = 1
    section = None
    gaps = []

    for raw in lines:
        s = (raw or "").strip()
        if not s:
            continue

        sm = SECTION_RE.match(s)
        if sm and is_section_header(s):
            if cur:
                blocks.append(cur)
                cur = None
            section = normalize_section(s)
            expected = 1
            continue

        m = QSTART.match(s)
        if m:
            num = int(m.group(1))
            if cur is None or num == expected or (expected < num <= expected + 10):
                if cur:
                    blocks.append(cur)
                if cur is not None and num != expected:
                    gaps.append((expected, num))
                cur = {"no": num, "section": section, "lines": [m.group(2)]}
                expected = num + 1
                continue

        if cur:
            cur["lines"].append(s)

    if cur:
        blocks.append(cur)

    questions = []
    for b in blocks:
        ls = b["lines"]
        ans_idx = None
        answer = None
        for i, l in enumerate(ls):
            am = ANS_LINE.search(l)
            if am:
                ans_idx = i
                raw_answer = am.group(1).strip()
                if raw_answer in ("正确", "错误"):
                    answer = raw_answer
                else:
                    answer = normalize_letters(raw_answer)
                break

        if ans_idx is None:
            stem = clean_stem(join_wrapped(ls))
            questions.append(_finish_question(
                b["no"], b["section"], stem, "", False, {},
                "原文未找到“正确答案”行"))
            continue

        stem_lines = []
        option_lines = []
        in_options = False
        for l in ls[:ans_idx]:
            if OPTSTART.match(l):
                in_options = True
            if in_options:
                option_lines.append(l)
            else:
                stem_lines.append(l)

        stem = join_wrapped(stem_lines)
        stem = strip_number_prefix(stem)

        # 同一行里混入的选项
        inline_start = find_inline_options_start(stem)
        if inline_start is not None:
            option_lines.insert(0, stem[inline_start:])
            stem = stem[:inline_start]

        stem = clean_stem(stem)
        stem = clean_answer_markers(stem, answer)
        options = extract_options(option_lines)
        judgment = (_is_judgment_options(options) and
                    answer in ("A", "B", "正确", "错误"))
        note = ""
        if judgment and answer in ("A", "B"):
            answer = "正确" if answer == "A" else "错误"

        questions.append(_finish_question(
            b["no"], b["section"], stem, answer, judgment, options, note))

    return questions, gaps


def parse_inline(lines, source=""):
    """解析答案写在题干括号/行尾的格式（2022 Word）。"""
    blocks = []
    cur = None
    expected = 1
    section = None
    gaps = []

    for raw in lines:
        s = (raw or "").strip()
        if not s:
            continue

        sm = SECTION_RE.match(s)
        if sm and is_section_header(s):
            if cur:
                blocks.append(cur)
                cur = None
            section = normalize_section(s)
            expected = 1
            continue

        m = QSTART.match(s)
        if m:
            num = int(m.group(1))
            if cur is None or num == expected or (expected < num <= expected + 10):
                if cur:
                    blocks.append(cur)
                if cur is not None and num != expected:
                    gaps.append((expected, num))
                cur = {"no": num, "section": section, "lines": [m.group(2)]}
                expected = num + 1
                continue

        if cur:
            cur["lines"].append(s)

    if cur:
        blocks.append(cur)

    questions = []
    for b in blocks:
        ls = b["lines"]
        opt_idx = None
        for i, l in enumerate(ls):
            if OPTSTART.match(l):
                opt_idx = i
                break

        if opt_idx is None:
            combined = join_wrapped(ls)
            cut = find_inline_options_start(combined)
            if cut is not None:
                stem_lines = [combined[:cut]]
                option_lines = [combined[cut:]]
            else:
                stem_lines = list(ls)
                option_lines = []
        else:
            stem_lines = list(ls[:opt_idx])
            option_lines = list(ls[opt_idx:])

        # 答案单独占一行，例如题干下一行只有一个“B”
        standalone = None
        for i in range(len(stem_lines) - 1, -1, -1):
            t = stem_lines[i].strip()
            if not t:
                continue
            if re.fullmatch(r"[A-Ga-g]{1,8}", t):
                standalone = t.upper()
                del stem_lines[i]
            break

        stem = join_wrapped(stem_lines)
        stem = strip_number_prefix(stem)

        answer = None
        judgment = False
        if standalone:
            answer = standalone
        else:
            answer, stem, judgment = extract_inline_answer(stem)

        if not answer:
            m = TRAILING_ANSWER.search(stem)
            if m:
                answer = m.group(1).upper()
                stem = stem[:m.start(1)].rstrip(" 。.、，,；;：:")

        stem = clean_stem(stem)
        if answer:
            stem = clean_answer_markers(stem, answer)
        options = extract_options(option_lines)

        note = ""
        if not answer:
            note = "原文未标注答案，请人工补充"

        questions.append(_finish_question(
            b["no"], b["section"], stem, answer, judgment, options, note))

    return questions, gaps


def _count_missing(questions):
    return sum(1 for q in questions if not q["answer"])


def parse_file(path, verbose=False):
    lines = read_source_lines(path)
    ans_count = sum(1 for l in lines if ANS_LINE.search(l))
    q_count = sum(1 for l in lines if QSTART.match(l.strip()))

    mode = "正确答案行"
    questions, gaps = parse_answerline(lines)
    if ans_count == 0 or ans_count < q_count * 0.3:
        mode = "题干内答案"
        questions, gaps = parse_inline(lines)
    else:
        # 少量块缺少“正确答案”时，尝试用题干内答案补一次
        inline_questions, _ = parse_inline(lines)
        if _count_missing(inline_questions) < _count_missing(questions):
            inline_by_no = {}
            for q in inline_questions:
                inline_by_no.setdefault(q["no"], q)
            for q in questions:
                if not q["answer"] and q["no"] in inline_by_no:
                    q["answer"] = inline_by_no[q["no"]]["answer"]
                    q["qtype"] = inline_by_no[q["no"]]["qtype"]
                    q["note"] = _append_note(q["note"], "答案由题干内标记补充")

    # 注意：不同章节的题号会从 1 重新开始，不能按题号去重。

    # 标记完全重复的题干
    groups = {}
    for q in questions:
        key = normalize_spaces(q["stem"]).lower()
        groups.setdefault(key, []).append(q)
    duplicate_count = 0
    for group in groups.values():
        if len(group) > 1:
            duplicate_count += len(group) - 1
            for q in group:
                q["note"] = _append_note(q["note"], "题干与其他题重复")

    bank = {
        "path": path,
        "source": os.path.basename(path),
        "mode": mode,
        "questions": questions,
        "gaps": gaps,
        "duplicates": duplicate_count,
    }
    if verbose:
        print("[解析] %s：%d 题（%s），缺答案 %d，重复 %d" %
              (bank["source"], len(questions), mode,
               _count_missing(questions), duplicate_count))
        if gaps:
            print("       题号跳号：%s" % ", ".join("%d->%d" % g for g in gaps[:10]))
    return bank


# ---------------------------------------------------------------------------
# Excel 输出
# ---------------------------------------------------------------------------

def _excel_safe(value, limit=32000):
    text = "" if value is None else str(value)
    text = EXCEL_BAD.sub("", text)
    if len(text) > limit:
        text = text[:limit] + "……"
    return text


def write_excel(banks, output_path, keep_missing=True):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ConvertError("缺少 openpyxl，请先执行：pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "题库"
    headers = ["题目", "答案", "题型", "原始题号", "来源", "备注"]
    ws.append(headers)

    row_count = 0
    for bank in banks:
        for q in bank["questions"]:
            if not q["answer"] and not keep_missing:
                continue
            ws.append([
                _excel_safe(q["stem"]),
                _excel_safe(q["answer"]),
                _excel_safe(q["qtype"]),
                q["no"],
                _excel_safe(bank["source"]),
                _excel_safe(q["note"]),
            ])
            row_count += 1

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [80, 14, 10, 10, 28, 34]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        row[1].alignment = Alignment(horizontal="center", vertical="top")
        row[2].alignment = Alignment(horizontal="center", vertical="top")
        row[3].alignment = Alignment(horizontal="center", vertical="top")
        row[4].alignment = Alignment(vertical="top")
        row[5].alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:F%d" % (row_count + 1)

    # 转换报告
    rs = wb.create_sheet("转换报告")
    rs.append(["文件", "题目数", "判断题", "单选题", "多选题", "其他",
               "缺答案", "重复", "解析方式"])
    for cell in rs[1]:
        cell.font = Font(bold=True)
    for bank in banks:
        counts = Counter(q["qtype"] for q in bank["questions"])
        rs.append([
            bank["source"], len(bank["questions"]),
            counts.get("判断题", 0), counts.get("单选题", 0),
            counts.get("多选题", 0), counts.get("其他", 0),
            _count_missing(bank["questions"]), bank["duplicates"],
            bank["mode"],
        ])
    for i, width in enumerate([36, 10, 10, 10, 10, 10, 10, 10, 14], start=1):
        rs.column_dimensions[get_column_letter(i)].width = width

    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir)
    try:
        wb.save(output_path)
    except PermissionError:
        raise ConvertError("无法写入 %s，请先关闭已在 Excel 中打开的该文件。" % output_path)
    return row_count


def default_output_path(input_path):
    base = os.path.splitext(os.path.basename(input_path))[0]
    return os.path.join(os.path.dirname(os.path.abspath(input_path)),
                        base + "_题库.xlsx")


# ---------------------------------------------------------------------------
# 命令行
# ---------------------------------------------------------------------------

def _expand_inputs(items):
    files = []
    for item in items:
        if os.path.isdir(item):
            for name in sorted(os.listdir(item)):
                path = os.path.join(item, name)
                if os.path.isfile(path) and os.path.splitext(name)[1].lower() in SUPPORTED_EXTS:
                    files.append(path)
        elif os.path.isfile(item):
            files.append(item)
        else:
            raise ConvertError("找不到文件：%s" % item)
    # 去重并保持顺序
    seen = set()
    result = []
    for f in files:
        key = os.path.abspath(f).lower()
        if key not in seen:
            seen.add(key)
            result.append(f)
    return result


def convert_cli(inputs, output=None, merge=False, keep_missing=True, preview=0):
    files = _expand_inputs(inputs)
    if not files:
        raise ConvertError("没有找到可转换的文件")

    banks = []
    for path in files:
        try:
            bank = parse_file(path, verbose=True)
        except ConvertError:
            raise
        except Exception as e:
            raise ConvertError("解析 %s 失败：%s" % (os.path.basename(path), e))
        banks.append(bank)

    if preview:
        print()
        for bank in banks:
            print("=" * 78)
            print(bank["source"], "（%d 题）" % len(bank["questions"]))
            for q in bank["questions"][:preview]:
                print("%4d [%s] %s  =>  %s" %
                      (q["no"], q["qtype"], q["stem"][:70], q["answer"]))
        print()

    if merge or len(files) == 1:
        out = output or (default_output_path(files[0]) if len(files) == 1
                         else os.path.join(os.path.dirname(os.path.abspath(files[0])),
                                           "题库合集.xlsx"))
        count = write_excel(banks, out, keep_missing=keep_missing)
        print("已生成：%s（%d 行）" % (out, count))
    else:
        if output:
            if os.path.isdir(output):
                out_dir = output
            elif os.path.splitext(output)[1].lower() == ".xlsx":
                # 传的是文件名时，把多个结果放到该文件所在目录
                out_dir = os.path.dirname(os.path.abspath(output))
            else:
                # 没有扩展名，按输出目录处理
                out_dir = output
        else:
            out_dir = os.path.dirname(os.path.abspath(files[0]))
        if not os.path.isdir(out_dir):
            os.makedirs(out_dir)
        for bank, path in zip(banks, files):
            out = os.path.join(out_dir,
                               os.path.splitext(os.path.basename(path))[0] + "_题库.xlsx")
            count = write_excel([bank], out, keep_missing=keep_missing)
            print("已生成：%s（%d 行）" % (out, count))

    total = sum(len(b["questions"]) for b in banks)
    missing = sum(_count_missing(b["questions"]) for b in banks)
    print("完成：%d 个文件，%d 题，缺答案 %d 题。" % (len(banks), total, missing))


# ---------------------------------------------------------------------------
# 图形界面
# ---------------------------------------------------------------------------

def _configure_gui_fonts(root, ttk_module):
    try:
        import tkinter.font as tkfont
        family = "Microsoft YaHei UI" if sys.platform == "win32" else "Segoe UI"
        available = set(tkfont.families(root))
        if family not in available:
            for fallback in ("Microsoft YaHei", "Segoe UI", "Arial"):
                if fallback in available:
                    family = fallback
                    break
        for name in ("TkDefaultFont", "TkTextFont", "TkFixedFont",
                     "TkMenuFont", "TkIconFont", "TkTooltipFont"):
            try:
                tkfont.nametofont(name).configure(family=family, size=10)
            except Exception:
                pass
        for name in ("TkHeadingFont", "TkCaptionFont"):
            try:
                tkfont.nametofont(name).configure(
                    family=family, size=10, weight="bold")
            except Exception:
                pass
        ttk_module.Style(root).configure(".", font=(family, 10))
        ttk_module.Style(root).configure(
            "TLabelframe.Label", font=(family, 10, "bold"))
    except Exception:
        pass


def launch_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    class App(object):
        def __init__(self, root):
            self.root = root
            self.files = []
            self.busy = False
            root.title("%s v%s - Word/PDF → Excel" % (APP_NAME, VERSION))
            root.geometry("1100x720")
            root.minsize(900, 600)

            top = ttk.Frame(root, padding=8)
            top.pack(fill="both", expand=False)

            file_box = ttk.LabelFrame(top, text="待转换文件（支持 .docx / .doc / .wps / .pdf）",
                                      padding=6)
            file_box.pack(fill="both", expand=False)

            list_frame = ttk.Frame(file_box)
            list_frame.pack(side="left", fill="both", expand=True)
            self.file_list = tk.Listbox(list_frame, height=6, selectmode="extended")
            self.file_list.pack(side="left", fill="both", expand=True)
            sb = ttk.Scrollbar(list_frame, orient="vertical",
                               command=self.file_list.yview)
            sb.pack(side="right", fill="y")
            self.file_list.config(yscrollcommand=sb.set)

            btns = ttk.Frame(file_box)
            btns.pack(side="right", padx=(8, 0), fill="y")
            ttk.Button(btns, text="添加文件", command=self.add_files).pack(fill="x", pady=2)
            ttk.Button(btns, text="添加文件夹", command=self.add_folder).pack(fill="x", pady=2)
            ttk.Button(btns, text="移除选中", command=self.remove_selected).pack(fill="x", pady=2)
            ttk.Button(btns, text="清空", command=self.clear_files).pack(fill="x", pady=2)

            opt = ttk.LabelFrame(top, text="输出设置", padding=6)
            opt.pack(fill="x", pady=(8, 0))
            self.merge_var = tk.BooleanVar(value=True)
            self.keep_missing_var = tk.BooleanVar(value=True)
            ttk.Checkbutton(opt, text="多个文件合并为一个 Excel",
                            variable=self.merge_var).grid(row=0, column=0, sticky="w")
            ttk.Checkbutton(opt, text="保留原文缺少答案的题目（答案留空）",
                            variable=self.keep_missing_var).grid(row=0, column=1,
                                                                sticky="w", padx=(18, 0))
            ttk.Label(opt, text="输出文件/目录：").grid(row=1, column=0,
                                                    sticky="w", pady=(6, 0))
            self.output_var = tk.StringVar()
            ttk.Entry(opt, textvariable=self.output_var).grid(
                row=1, column=1, sticky="we", pady=(6, 0))
            ttk.Button(opt, text="浏览…", command=self.browse_output).grid(
                row=1, column=2, padx=(6, 0), pady=(6, 0))
            opt.columnconfigure(1, weight=1)

            actions = ttk.Frame(top)
            actions.pack(fill="x", pady=(8, 0))
            self.preview_btn = ttk.Button(actions, text="预览选中文件",
                                          command=self.preview_selected)
            self.preview_btn.pack(side="left")
            self.convert_btn = ttk.Button(actions, text="开始转换",
                                          command=self.start_convert)
            self.convert_btn.pack(side="left", padx=(8, 0))
            self.status_var = tk.StringVar(value="就绪")
            ttk.Label(actions, textvariable=self.status_var).pack(side="left", padx=(16, 0))
            self.progress = ttk.Progressbar(actions, mode="indeterminate", length=160)
            self.progress.pack(side="right")

            table_box = ttk.LabelFrame(root, text="预览", padding=6)
            table_box.pack(fill="both", expand=True, padx=8, pady=(0, 8))
            columns = ("no", "type", "stem", "answer", "note")
            self.tree = ttk.Treeview(table_box, columns=columns, show="headings")
            headings = {"no": "题号", "type": "题型", "stem": "题目",
                        "answer": "答案", "note": "备注"}
            widths = {"no": 50, "type": 70, "stem": 680, "answer": 90, "note": 160}
            for col in columns:
                self.tree.heading(col, text=headings[col])
                self.tree.column(col, width=widths[col], anchor="w")
            vsb = ttk.Scrollbar(table_box, orient="vertical", command=self.tree.yview)
            hsb = ttk.Scrollbar(table_box, orient="horizontal", command=self.tree.xview)
            self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
            self.tree.grid(row=0, column=0, sticky="nsew")
            vsb.grid(row=0, column=1, sticky="ns")
            hsb.grid(row=1, column=0, sticky="we")
            table_box.rowconfigure(0, weight=1)
            table_box.columnconfigure(0, weight=1)

            self.log = tk.Text(root, height=4, state="disabled")
            self.log.pack(fill="x", padx=8, pady=(0, 8))

        # ---- 文件操作 ----
        def add_files(self):
            paths = filedialog.askopenfilenames(
                title="选择题库文件",
                filetypes=[("题库文件", "*.docx *.doc *.wps *.pdf"),
                           ("Word 文档", "*.docx *.doc *.wps"),
                           ("PDF 文件", "*.pdf"),
                           ("所有文件", "*.*")])
            self._add(paths)

        def add_folder(self):
            folder = filedialog.askdirectory(title="选择题库文件夹")
            if not folder:
                return
            paths = []
            for name in sorted(os.listdir(folder)):
                path = os.path.join(folder, name)
                if os.path.isfile(path) and os.path.splitext(name)[1].lower() in SUPPORTED_EXTS:
                    paths.append(path)
            self._add(paths)

        def _add(self, paths):
            for path in paths:
                if not path:
                    continue
                ext = os.path.splitext(path)[1].lower()
                if ext not in SUPPORTED_EXTS:
                    continue
                if any(os.path.abspath(path).lower() == os.path.abspath(x).lower()
                       for x in self.files):
                    continue
                self.files.append(path)
                self.file_list.insert("end", "%s  (%s)" %
                                      (os.path.basename(path),
                                       os.path.getsize(path) // 1024 + 1))
            if self.files and not self.output_var.get():
                first = self.files[0]
                self.output_var.set(os.path.join(
                    os.path.dirname(os.path.abspath(first)), "题库合集.xlsx"))

        def remove_selected(self):
            for index in reversed(self.file_list.curselection()):
                del self.files[index]
                self.file_list.delete(index)

        def clear_files(self):
            self.files = []
            self.file_list.delete(0, "end")
            self.tree.delete(*self.tree.get_children())

        def browse_output(self):
            if self.merge_var.get() or len(self.files) <= 1:
                path = filedialog.asksaveasfilename(
                    title="保存 Excel", defaultextension=".xlsx",
                    filetypes=[("Excel 文件", "*.xlsx")])
                if path:
                    self.output_var.set(path)
            else:
                folder = filedialog.askdirectory(title="选择输出文件夹")
                if folder:
                    self.output_var.set(folder)

        # ---- 日志 ----
        def log_line(self, text):
            self.log.config(state="normal")
            self.log.insert("end", text + "\n")
            self.log.see("end")
            self.log.config(state="disabled")

        def set_busy(self, busy, text=None):
            self.busy = busy
            state = "disabled" if busy else "normal"
            self.preview_btn.config(state=state)
            self.convert_btn.config(state=state)
            if busy:
                self.progress.start(12)
            else:
                self.progress.stop()
            if text is not None:
                self.status_var.set(text)

        # ---- 预览 ----
        def preview_selected(self):
            if not self.files:
                messagebox.showwarning("提示", "请先添加文件")
                return
            selection = self.file_list.curselection()
            index = selection[0] if selection else 0
            path = self.files[index]
            self.set_busy(True, "正在解析预览…")
            self._run_async(lambda: parse_file(path, verbose=False),
                            self._preview_done)

        def _preview_done(self, bank=None, error=None):
            self.set_busy(False, "就绪")
            if error:
                messagebox.showerror("解析失败", error)
                return
            self.tree.delete(*self.tree.get_children())
            for q in bank["questions"][:500]:
                self.tree.insert("", "end", values=(
                    q["no"], q["qtype"], q["stem"][:300], q["answer"], q["note"]))
            counts = Counter(q["qtype"] for q in bank["questions"])
            self.status_var.set(
                "%s：%d 题（判断 %d / 单选 %d / 多选 %d / 其他 %d），缺答案 %d" % (
                    bank["source"], len(bank["questions"]),
                    counts.get("判断题", 0), counts.get("单选题", 0),
                    counts.get("多选题", 0), counts.get("其他", 0),
                    _count_missing(bank["questions"])))
            self.log_line("解析：%s，方式：%s" % (bank["source"], bank["mode"]))
            if bank["gaps"]:
                self.log_line("题号跳号：" + ", ".join(
                    "%d->%d" % g for g in bank["gaps"][:10]))

        # ---- 转换 ----
        def start_convert(self):
            if not self.files:
                messagebox.showwarning("提示", "请先添加文件")
                return
            output = self.output_var.get().strip()
            merge = self.merge_var.get() or len(self.files) == 1
            if merge:
                if not output:
                    messagebox.showwarning("提示", "请选择输出 Excel 文件")
                    return
                if not output.lower().endswith(".xlsx"):
                    output += ".xlsx"
                    self.output_var.set(output)
            self.set_busy(True, "正在转换…")
            keep_missing = self.keep_missing_var.get()
            files = list(self.files)

            def job():
                banks = [parse_file(p, verbose=False) for p in files]
                if merge:
                    count = write_excel(banks, output, keep_missing=keep_missing)
                    return banks, [output], count
                out_dir = output if os.path.isdir(output) or not os.path.splitext(output)[1] \
                    else os.path.dirname(os.path.abspath(output))
                if not os.path.isdir(out_dir):
                    os.makedirs(out_dir)
                outputs = []
                total = 0
                for bank, path in zip(banks, files):
                    out = os.path.join(
                        out_dir,
                        os.path.splitext(os.path.basename(path))[0] + "_题库.xlsx")
                    total += write_excel([bank], out, keep_missing=keep_missing)
                    outputs.append(out)
                return banks, outputs, total

            self._run_async(job, self._convert_done)

        def _convert_done(self, result=None, error=None):
            self.set_busy(False, "就绪")
            if error:
                messagebox.showerror("转换失败", error)
                self.log_line("失败：" + error)
                return
            banks, outputs, count = result
            total = sum(len(b["questions"]) for b in banks)
            missing = sum(_count_missing(b["questions"]) for b in banks)
            self.status_var.set("完成：%d 题，缺答案 %d 题" % (total, missing))
            for out in outputs:
                self.log_line("已生成：" + out)
            messagebox.showinfo(
                "转换完成",
                "共 %d 个文件、%d 题，缺答案 %d 题。\n\n输出：\n%s" %
                (len(banks), total, missing, "\n".join(outputs)))

        # ---- 异步执行 ----
        def _run_async(self, func, done):
            def worker():
                try:
                    result = func()
                except Exception as e:
                    err = str(e)
                    if not isinstance(e, ConvertError):
                        err = err + "\n\n" + traceback.format_exc()
                    self.root.after(0, lambda: done(None, err))
                else:
                    self.root.after(0, lambda: done(result, None))
            threading.Thread(target=worker, daemon=True).start()

    root = tk.Tk()
    _configure_gui_fonts(root, ttk)
    App(root)
    root.mainloop()


# ---------------------------------------------------------------------------
# 一体化版内部调用接口
# ---------------------------------------------------------------------------

_TEMP_DIRS = []
_CACHE = {}
_CACHE_BANKS = {}


def write_simple_excel(bank, output_path):
    """一体化版内部使用：只写“题目/答案”两列，避免主界面出现多余列。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise ConvertError("缺少 openpyxl，请先执行：pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "题库"
    ws.append(["题目", "答案"])
    for q in bank["questions"]:
        ws.append([_excel_safe(q["stem"]), _excel_safe(q["answer"])])

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions["A"].width = 100
    ws.column_dimensions["B"].width = 18
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        row[1].alignment = Alignment(horizontal="center", vertical="top")
    ws.freeze_panes = "A2"
    wb.save(output_path)


def import_tiku_file(source_path):
    """把 Word/PDF 题库转换成临时 Excel，返回 (xlsx路径, 解析结果)。

    同一个文件（路径、大小、修改时间都相同）重复加载时直接复用临时文件。
    临时目录在程序退出时由 cleanup_temp_files() 清理。
    """
    source_path = os.path.abspath(source_path)
    try:
        st = os.stat(source_path)
        key = (source_path, st.st_mtime_ns, st.st_size)
    except Exception:
        key = (source_path, 0, 0)

    cached = _CACHE.get(key)
    if cached and os.path.exists(cached):
        return cached, _CACHE_BANKS.get(key)

    bank = parse_file(source_path, verbose=False)
    temp_dir = tempfile.mkdtemp(prefix="tiku_import_")
    out_path = os.path.join(temp_dir, "converted.xlsx")
    try:
        write_simple_excel(bank, out_path)
    except Exception:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise

    _TEMP_DIRS.append(temp_dir)
    _CACHE[key] = out_path
    _CACHE_BANKS[key] = bank
    return out_path, bank


def cleanup_temp_files():
    """清理一体化版解析 Word/PDF 时生成的临时 Excel。"""
    for folder in list(_TEMP_DIRS):
        shutil.rmtree(folder, ignore_errors=True)
    del _TEMP_DIRS[:]
    _CACHE.clear()
    _CACHE_BANKS.clear()


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------

def main(argv=None):
    parser = argparse.ArgumentParser(
        description="把 Word/PDF 题库转换为答题工具可直接读取的 Excel")
    parser.add_argument("inputs", nargs="*", help="题库文件或文件夹")
    parser.add_argument("-o", "--output", help="输出 Excel 或输出目录")
    parser.add_argument("--merge", action="store_true", help="多个文件合并为一个 Excel")
    parser.add_argument("--drop-missing", action="store_true",
                        help="不输出缺少答案的题目")
    parser.add_argument("--preview", type=int, default=0,
                        help="在终端预览每个文件的前 N 题")
    parser.add_argument("--gui", action="store_true", help="打开图形界面")
    args = parser.parse_args(argv)

    if args.gui or not args.inputs:
        launch_gui()
        return 0

    try:
        convert_cli(args.inputs, output=args.output, merge=args.merge,
                    keep_missing=not args.drop_missing, preview=args.preview)
    except ConvertError as e:
        print("错误：%s" % e, file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print("已取消", file=sys.stderr)
        return 130
    return 0


if __name__ == "__main__":
    sys.exit(main())
